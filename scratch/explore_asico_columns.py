import openpyxl
from pathlib import Path
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = Path(r"C:\Users\Riew\Desktop\Job_NBV\KRC\krc-erp\docs\รายการชิ้นงานลูกค้า ASICO 9324 (1).xlsx")
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['รายการชิ้นงาน โรง 1']

row3 = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
row4 = next(ws.iter_rows(min_row=4, max_row=4, values_only=True))

print("Row 3 values:")
for idx, val in enumerate(row3):
    print(f"  Col {idx}: {val}")

print("\nRow 4 values:")
for idx, val in enumerate(row4):
    print(f"  Col {idx}: {val}")
