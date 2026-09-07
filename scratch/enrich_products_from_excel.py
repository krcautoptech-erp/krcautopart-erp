import openpyxl
import os
import re
import json
import sys

# Force UTF-8 encoding for stdout
sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"C:\Users\Riew\Desktop\Job_NBV\KRC\krc-erp\docs\รายการชิ้นงานลูกค้า ASICO 9324 (1).xlsx"
output_json_path = r"C:\Users\Riew\Desktop\Job_NBV\KRC\krc-erp\scratch\temp_excel_data.json"

def normalize_key(part):
    if not part:
        return ""
    return re.sub(r'[^A-Z0-9]', '', str(part).upper())

def main():
    print(f"Parsing Excel file: {os.path.basename(excel_path)}...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    sheet_name = 'รายการชิ้นงาน โรง 1'
    if sheet_name not in wb.sheetnames:
        print(f"Error: Sheet '{sheet_name}' not found. Available sheets: {wb.sheetnames}")
        return
        
    ws = wb[sheet_name]
    
    # We load all data from Row 5 onwards
    rows = list(ws.iter_rows(min_row=5, values_only=True))
    excel_records = {}
    
    for idx, r in enumerate(rows):
        part_no = r[4] # Col 4 is หมายเลขชิ้นส่วน
        if not part_no:
            continue
            
        part_no_str = str(part_no).strip()
        if not part_no_str or part_no_str == 'Part Namber':
            continue
            
        # Extract fields
        std_no = str(r[1]).strip() if r[1] is not None else None
        material = str(r[3]).strip() if r[3] is not None else None
        model = str(r[6]).strip() if r[6] is not None else None
        
        if model:
            model = model.strip()
            
        norm_key = normalize_key(part_no_str)
        
        # Save record
        if norm_key not in excel_records:
            excel_records[norm_key] = {
                'part_number_orig': part_no_str,
                'std_no': std_no,
                'material': material,
                'model': model
            }
        else:
            rec = excel_records[norm_key]
            if not rec['std_no'] and std_no:
                rec['std_no'] = std_no
            if not rec['material'] and material:
                rec['material'] = material
            if not rec['model'] and model:
                rec['model'] = model

    print(f"Loaded {len(excel_records)} unique part numbers from Excel sheet. Exporting to JSON...")
    with open(output_json_path, "w", encoding="utf-8") as f:
        json.dump(excel_records, f, ensure_ascii=False, indent=2)
    print("Export complete.")

if __name__ == "__main__":
    main()
