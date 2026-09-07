import sys
from pathlib import Path
from openpyxl import load_workbook


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    path = Path(r"C:\Users\Riew\Desktop\Job_NBV\KRC\krc-erp\docs\รายการชิ้นงานลูกค้า ASICO 9324 (1).xlsx")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["รายการชิ้นงาน โรง 1"]
    for r in range(1, 30):
        vals = [ws.cell(r, c).value for c in range(1, 11)]
        print(r, vals)


if __name__ == "__main__":
    main()
