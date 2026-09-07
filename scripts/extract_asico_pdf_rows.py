import json
import sys
from pathlib import Path
from openpyxl import load_workbook


SOURCE_PATH = Path(r"C:\Users\Riew\Desktop\Job_NBV\KRC\krc-erp\docs\รายการชิ้นงานลูกค้า ASICO 9324 (1).xlsx")

SPECS = [
    {
        "source_sheet": "รายการชิ้นงาน โรง 1",
        "output_sheet": "รายการชิ้นงาน โรง 1",
        "item_start": 1,
        "item_end": 177,
        "headers": ["NO.", "STD", "จำนวนแผ่น", "MATERIAL", "Part Number", "Part Name", "MODEL", "ชิ้นงานต่อแผ่น"],
        "columns": [1, 2, 3, 4, 5, 6, 7, 9],
    },
    {
        "source_sheet": "รายการชิ้นงานโรง 2 ",
        "output_sheet": "รายการชิ้นงาน โรง 2",
        "item_start": 1,
        "item_end": 49,
        "headers": ["NO.", "SPEC", "จำนวนแผ่น", "MATERIAL", "Part Number", "Part Name", "MODEL", "PLATING / ชุบ", "ชิ้นงานต่อแผ่น"],
        "columns": [1, 2, 3, 4, 5, 6, 7, 8, 10],
    },
]


def cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def build_rows(ws, item_start: int, item_end: int, columns: list[int]) -> list[list[str]]:
    rows: list[list[str]] = []
    all_rows = list(ws.iter_rows(min_row=5, values_only=True))
    for offset, row_values in enumerate(all_rows):
        row_index = offset + 5
        item_no = row_values[0]
        if isinstance(item_no, (int, float)):
            item_no = int(item_no)
            if item_no > item_end:
                break
            if item_no >= item_start:
                values = [cell_text(row_values[c - 1]) for c in columns]
                suffix_idx = offset + 2
                suffix_row = all_rows[suffix_idx] if suffix_idx < len(all_rows) else None
                model_suffix = cell_text(suffix_row[6] if suffix_row and len(suffix_row) > 6 else "")
                if model_suffix:
                    values[6] = f"{values[6]}\n{model_suffix}" if values[6] else model_suffix
                rows.append(values)
    return rows


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    wb = load_workbook(SOURCE_PATH, read_only=True, data_only=True)
    payload = {"sheets": []}
    for spec in SPECS:
        ws = wb[spec["source_sheet"]]
        payload["sheets"].append(
            {
                "name": spec["output_sheet"],
                "headers": spec["headers"],
                "rows": build_rows(ws, spec["item_start"], spec["item_end"], spec["columns"]),
            }
        )

    if len(sys.argv) > 1:
        output_path = Path(sys.argv[1])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(output_path)
    else:
        print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
