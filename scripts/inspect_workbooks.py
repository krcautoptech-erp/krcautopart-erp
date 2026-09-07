import sys
from pathlib import Path
from openpyxl import load_workbook


def describe(path: Path) -> None:
    print(f"FILE\t{path.name}")
    wb = load_workbook(path, read_only=True, data_only=True)
    print("SHEETS\t" + ", ".join(wb.sheetnames))
    for sheet_name in wb.sheetnames[:3]:
        ws = wb[sheet_name]
        print(f"SHEET\t{sheet_name}\trows={ws.max_row}\tcols={ws.max_column}")
        for row in ws.iter_rows(min_row=1, max_row=min(8, ws.max_row), values_only=True):
            print("ROW\t" + "\t".join("" if v is None else str(v) for v in row[:10]))
    print()


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    base = Path(r"C:\Users\Riew\Desktop\Job_NBV\KRC\krc-erp\docs")
    for path in sorted(base.glob("*.xlsx")):
        describe(path)


if __name__ == "__main__":
    main()
